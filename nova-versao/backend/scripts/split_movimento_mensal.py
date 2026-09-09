"""Divide EXITO Entradas/Saídas acumulado em um arquivo por YYYY-MM (Data Emissão)."""
from __future__ import annotations

import argparse
import calendar
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402

from app.extract.classify import parse_br_number, scan_cnpj, scan_razao  # noqa: E402
from app.extract.parse_movimento import _header_map, _is_detail_code, _skip_row  # noqa: E402
from app.extract.workbook import load_workbook  # noqa: E402

DATE_LABELS = ("data emissao", "data emissão", "dt emissao", "emissao")


def _norm(s: str) -> str:
    t = (s or "").strip().lower()
    trans = str.maketrans("áàâãéêíóôõúç", "aaaaeeiooouc")
    t = t.translate(trans)
    return re.sub(r"\s+", " ", t)


def _excel_serial_to_dt(value: float) -> datetime | None:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return None
    if n < 20000 or n > 80000:
        return None
    return datetime(1899, 12, 30) + timedelta(days=n)


def parse_date(cell: str) -> datetime | None:
    raw = (cell or "").strip()
    if not raw:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw[:10], fmt)
        except ValueError:
            continue
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", raw)
    if m:
        return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    try:
        n = float(raw.replace(",", "."))
    except ValueError:
        return None
    if n == int(n) and 20000 < n < 80000:
        return _excel_serial_to_dt(n)
    return None


def _date_col(header_row: list[str], mapping: dict[str, int]) -> int | None:
    labels = [_norm(c) for c in header_row]
    for i, lab in enumerate(labels):
        if any(k in lab for k in DATE_LABELS):
            return i
    # fallback típico EXITO: col 1 = Data Emissão
    if mapping.get("codigo") == 0:
        return 1
    return None


def _valor_idx(mapping: dict[str, int]) -> int:
    return int(mapping.get("valor", 0))


def split_file(src: Path, out_dir: Path, tipo: str) -> list[dict]:
    grid = load_workbook(src)
    header = None
    header_row_i = -1
    for i, row in enumerate(grid.rows[:40]):
        found = _header_map(row)
        if found:
            header = found
            header_row_i = i
            break
    if not header:
        raise SystemExit(f"Cabeçalho de movimento não encontrado: {src.name}")

    date_i = _date_col(grid.rows[header_row_i], header)
    if date_i is None:
        raise SystemExit(f"Coluna Data Emissão não encontrada: {src.name}")

    prefix = grid.rows[: header_row_i + 1]
    by_month: dict[str, list[list[str]]] = defaultdict(list)
    skipped = 0
    for row in grid.rows[header_row_i + 1 :]:
        if not row:
            continue
        joined = " ".join(str(c) for c in row).lower()
        if "total geral" in joined:
            continue
        codigo = row[header.get("codigo", 0)] if header.get("codigo", 0) < len(row) else ""
        if not _is_detail_code(str(codigo or "")):
            continue
        if _skip_row(row):
            continue
        cell = row[date_i] if date_i < len(row) else ""
        dt = parse_date(str(cell or ""))
        if not dt:
            skipped += 1
            continue
        key = dt.strftime("%Y-%m")
        by_month[key].append(list(row))

    src_cnpj = scan_cnpj(grid)
    src_razao = scan_razao(grid)

    out_dir.mkdir(parents=True, exist_ok=True)
    kind_name = "Entradas" if tipo == "entradas" else "Saidas"
    valor_i = _valor_idx(header)
    reports: list[dict] = []

    for competencia in sorted(by_month):
        lines = by_month[competencia]
        year, month = competencia.split("-")
        last = calendar.monthrange(int(year), int(month))[1]
        period = f"Período: 01/{month}/{year} até {last:02d}/{month}/{year}"
        rows_out: list[list[str]] = []
        for r in prefix:
            copy = list(r)
            joined = _norm(" ".join(str(c) for c in copy))
            if "periodo" in joined or "competencia" in joined:
                copy = [period]
                rows_out.append(copy)
                continue
            rows_out.append(copy)
        if src_cnpj:
            cnpj_row = None
            for r in rows_out:
                if "cnpj" in _norm(" ".join(str(c) for c in r)):
                    cnpj_row = r
                    break
            if cnpj_row is None:
                rows_out.insert(1, ["CNPJ:", src_cnpj])
            else:
                cnpj_row[:] = ["CNPJ:", src_cnpj]
        if src_razao and rows_out:
            first = _norm(str(rows_out[0][0] if rows_out[0] else ""))
            if "jpg" not in first and "periodo" in first:
                rows_out.insert(0, [src_razao])
        rows_out.extend(lines)
        soma = 0.0
        for ln in lines:
            if valor_i < len(ln):
                soma += float(parse_br_number(str(ln[valor_i] or "0")) or 0)
        total_row = [""] * max(valor_i + 1, 8)
        total_row[0] = "Total Geral"
        total_row[valor_i] = f"{soma:.2f}".replace(".", ",")
        rows_out.append(total_row)

        dest = out_dir / f"{kind_name} {month}-{year}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = kind_name
        for row in rows_out:
            ws.append(row)
        wb.save(dest)
        reports.append(
            {
                "file": dest.name,
                "competencia": competencia,
                "lines": len(lines),
                "soma": round(soma, 2),
                "path": str(dest),
            }
        )
        print(f"OK {dest.name} linhas={len(lines)} soma={soma:.2f}")

    if skipped:
        print(f"WARN {src.name}: {skipped} linhas sem Data Emissão")
    return reports


def detect_tipo(name: str) -> str:
    low = name.lower()
    if "saida" in low or "saída" in low:
        return "saidas"
    return "entradas"


def main() -> int:
    p = argparse.ArgumentParser(description="Split movimento EXITO acumulado em arquivos mensais")
    p.add_argument("src", type=Path, help="Arquivo .xls acumulado")
    p.add_argument("--out", type=Path, required=True, help="Pasta de saída")
    p.add_argument("--tipo", choices=("entradas", "saidas"), default="")
    p.add_argument("--unidade", default="", help="Informativo (não grava Postgres)")
    args = p.parse_args()
    src = args.src
    if not src.exists():
        print("ERR arquivo ausente", src)
        return 1
    tipo = args.tipo or detect_tipo(src.name)
    print(f"SPLIT {src.name} tipo={tipo} unidade={args.unidade or '-'}")
    split_file(src, args.out, tipo)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
