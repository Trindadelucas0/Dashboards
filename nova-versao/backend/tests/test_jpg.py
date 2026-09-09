"""JPG — catálogo, unidades, split mensal e tabela ICMS/IPI."""
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.companies import COMPANY_BY_ID, KEEP_USERNAMES
from app.extract.classify import resolve_company, unit_from_filename
from app.extract.parse_impostos import composicao_from_apuracao
from app.extract.pipeline import classify_and_extract
from app.routers.companies import aggregate_fiscal_packs, catalog_units_payload
from app.extract.aggregate import merge_saidas
from app.extract.parse_movimento import ExtractedMovimento, Line
from scripts.seed_jpg_lannic import (
    ALIQUOTA,
    BASE_MEMORIA,
    DAS,
    PARTILHA,
    TOTAL_SAIDAS,
    build_lannic_agosto_2026,
    merge_pgdas_into_pack,
)

ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ENT = ROOT / "fixtures" / "jpg-padrao" / "Entradas 01-2026.xlsx"


def test_jpg_catalog_units():
    reg = COMPANY_BY_ID["jpg"]
    assert reg.username == "jpg"
    assert reg.cnpj == "21051983000165"
    keys = [u.key for u in reg.units]
    assert keys == ["sede", "asa_sul", "pr", "sp", "mg", "lannic"]
    assert "jpg" in KEEP_USERNAMES
    payload = catalog_units_payload("jpg", None)
    assert payload[0] == {"key": "todas", "label": "Todas as unidades"}
    assert {u["key"] for u in payload} >= {"todas", "sede", "pr", "mg", "sp", "asa_sul", "lannic"}
    assert "matriz" not in {u["key"] for u in payload}
    lannic = next(u for u in reg.units if u.key == "lannic")
    assert lannic.cnpj == "48285395000142"
    assert lannic.label == "LANNIC Dermocosméticos"


def test_jpg_filename_unit():
    assert unit_from_filename("ipi filial pr 01 a 08.xls") == "pr"
    assert unit_from_filename("icms filial asa sul 01 a 08.xls") == "asa_sul"
    assert unit_from_filename("ipi filial mg 01 a 08.xls") == "mg"
    assert unit_from_filename("pgdas lannic 08-2026.pdf") == "lannic"


def test_jpg_lannic_cnpj_resolve():
    company, unit = resolve_company("48285395000142", "LANNIC DERMOCOSMETICOS", "x.xls")
    assert company and company.id == "jpg"
    assert unit == "lannic"
    company, unit = resolve_company("", "LANNIC Dermocosméticos", "pgdas.xls")
    assert company and company.id == "jpg"
    assert unit == "lannic"


def test_jpg_lannic_das_composicao_and_golden():
    pack = build_lannic_agosto_2026()
    ap = pack["apuracao"]
    assert ap["fonte"] == "pgdas_simples_nacional"
    assert ap["das"]["aRecolher"] == DAS
    assert ap["das"]["apurado"] == DAS
    assert ap["das"]["aliquota"] == ALIQUOTA
    assert "icms" not in ap
    assert "irpj" not in ap
    assert "csll" not in ap
    comp = composicao_from_apuracao(ap)
    assert comp == [{"label": "Simples Nacional", "valor": DAS}]
    assert pack["receitaBruta"] == BASE_MEMORIA
    assert pack["cfopSaidasTotal"] == BASE_MEMORIA
    sn = pack["memoriaSimples"]
    assert sn["totalSaidas"] == TOTAL_SAIDAS
    assert sn["baseMemoria"] == BASE_MEMORIA
    assert "rpaPgdas" not in sn
    assert "diferencaBases" not in sn
    assert sn["partilha"] == PARTILHA
    assert pytest.approx(BASE_MEMORIA * ALIQUOTA / 100.0, abs=0.02) == DAS
    assert pack["deducoes"] == DAS


def test_jpg_lannic_merge_saidas_keeps_pgdas_receita():
    pack = build_lannic_agosto_2026()
    mov = ExtractedMovimento(
        tipo="saidas",
        company="LANNIC",
        cnpj="48285395000142",
        period="Período: 01/08/2026 até 31/08/2026",
        total_geral=893349.52,
        lines=[
            Line("1", "1", "1", "CLIENTE", "11111111000111", "DF", "6-108", 893349.52),
        ],
    )
    merged = merge_saidas(pack, mov)
    assert merged["cfopSaidasTotal"] == pytest.approx(893349.52, abs=0.02)
    assert merged["receitaBruta"] == BASE_MEMORIA
    assert merged["apuracao"]["das"]["aRecolher"] == DAS
    assert merged["memoriaSimples"]["baseMemoria"] == BASE_MEMORIA
    assert merged["hasMovimentacao"] is True


def test_jpg_lannic_seed_merge_keeps_movimento():
    existing = {
        "hasMovimentacao": True,
        "totalCompras": 451739.03,
        "cfopSaidasTotal": 893349.52,
        "nfsEntradas": 237,
        "nfsSaidas": 241,
        "receitaBruta": 893349.52,
        "apuracao": {"ipi": {"aRecolher": 1}},
    }
    merged = merge_pgdas_into_pack(existing, build_lannic_agosto_2026())
    assert merged["hasMovimentacao"] is True
    assert merged["totalCompras"] == pytest.approx(451739.03, abs=0.02)
    assert merged["cfopSaidasTotal"] == pytest.approx(893349.52, abs=0.02)
    assert merged["receitaBruta"] == BASE_MEMORIA
    assert merged["apuracao"]["das"]["aRecolher"] == DAS
    assert merged["apuracao"]["fonte"] == "pgdas_simples_nacional"
    assert merged["nfsEntradas"] == 237


def test_split_and_probe_movimento_lannic(tmp_path: Path):
    from scripts.split_movimento_mensal import split_file

    wb = Workbook()
    ws = wb.active
    ws.title = "Saidas"
    ws.append(["144 - LANNIC DERMOCOSMETICOS LTDA"])
    ws.append(["CNPJ:", "48.285.395/0001-42"])
    ws.append(["Período: 01/08/2026 até 31/08/2026"])
    ws.append(
        [
            "Código",
            "Data Emissão",
            "Nota",
            "Série",
            "Esp",
            "X",
            "Y",
            "Z",
            "W",
            "Q",
            "Cliente",
            "A",
            "CNPJ/CPF",
            "IE",
            "B",
            "C",
            "CFOP",
            "AC",
            "UF",
            "Valor Contábil",
        ]
    )
    ws.append(
        [1, "15/08/2026", "200", "1", "", "", "", "", "", "", "CLI A", "", "11111111000111", "", "", "", "5.102", "", "DF", "557733,52"]
    )
    src = tmp_path / "144-Saidas 01-2026 a 08-2026.xlsx"
    wb.save(src)
    out = tmp_path / "out"
    reports = split_file(src, out, "saidas")
    assert {r["competencia"] for r in reports} == {"2026-08"}
    ago = next(p for p in out.glob("*.xlsx") if "08-2026" in p.name)
    result = classify_and_extract(ago)
    assert result["tipo"] == "saidas"
    assert result["company_id"] == "jpg"
    assert result["unidade"] == "lannic"
    assert result["competencia"] == "2026-08"
    assert not result["errors"]
    assert abs(result["meta"]["delta"]) < 0.02
    assert result["pack_patch"]["cfopSaidasTotal"] == pytest.approx(557733.52, abs=0.02)


SPLIT_LANNIC = ROOT.parent / "pasta temporaria" / "_split" / "144" / "lannic"


@pytest.mark.skipif(not (SPLIT_LANNIC / "Saidas 08-2026.xlsx").exists(), reason="Split LANNIC ago/2026 ausente")
def test_jpg_lannic_agosto_split_golden():
    ent = classify_and_extract(SPLIT_LANNIC / "Entradas 08-2026.xlsx")
    sai = classify_and_extract(SPLIT_LANNIC / "Saidas 08-2026.xlsx")
    assert ent["company_id"] == sai["company_id"] == "jpg"
    assert ent["unidade"] == sai["unidade"] == "lannic"
    assert ent["tipo"] == "entradas"
    assert sai["tipo"] == "saidas"
    assert not ent["errors"] and not sai["errors"]
    assert abs(ent["meta"]["delta"]) < 0.02
    assert abs(sai["meta"]["delta"]) < 0.02
    assert ent["pack_patch"]["totalCompras"] == pytest.approx(451739.03, abs=0.02)
    assert sai["pack_patch"]["cfopSaidasTotal"] == pytest.approx(893349.52, abs=0.02)


def test_jpg_impostos_table_parts(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Empresa",
            "Filial",
            "Mês",
            "ICMS Crédito",
            "ICMS Débito",
            "ICMS a Recolher",
            "IPI Crédito",
            "IPI Débito",
            "IPI a Recolher",
        ]
    )
    ws.append(["JPG", "Filial PR", "Janeiro", 0, 0, 0, 10, 50, 40])
    ws.append(["JPG", "Filial PR", "Fevereiro", 0, 0, 0, 5, 20, 15])
    ws.append(["JPG", "Filial MG", "Janeiro", 0, 0, 0, 1, 2, 1])
    path = tmp_path / "ipi filial pr 01 a 08.xlsx"
    wb.save(path)
    result = classify_and_extract(path)
    assert result["company_id"] == "jpg"
    assert result["tipo"] == "impostos_mensal"
    parts = result.get("parts") or []
    assert len(parts) == 2
    assert all(p["unidade"] == "pr" for p in parts)
    jan = next(p for p in parts if p["competencia"] == "2026-01")
    assert jan["pack_patch"]["apuracao"]["ipi"]["aRecolher"] == 40
    assert not result.get("errors")


def test_jpg_ipi_table_saldo_credor(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Empresa",
            "Filial",
            "Mês",
            "ICMS Crédito",
            "ICMS Débito",
            "ICMS a Recolher",
            "IPI Crédito",
            "IPI Débito",
            "IPI a Recolher",
        ]
    )
    ws.append(["JPG", "Filial PR", "Agosto", 0, 0, 0, 324978.79, 323063.81, 0])
    path = tmp_path / "ipi filial pr 01 a 08.xlsx"
    wb.save(path)
    result = classify_and_extract(path)
    parts = result.get("parts") or []
    ago = next(p for p in parts if p["competencia"] == "2026-08")
    ipi = ago["pack_patch"]["apuracao"]["ipi"]
    assert ipi["aRecolher"] == pytest.approx(-1914.98, abs=0.02)
    assert ipi["credito"] == pytest.approx(324978.79, abs=0.02)


def test_jpg_demonstrativo_ipi_saldo_credor():
    from app.extract.parse_impostos import apuracao_patch_from_demo, parse_demonstrativo_ipi
    from app.extract.workbook import WorkbookGrid

    def row(label, valor):
        cells = [""] * 13
        cells[0] = label
        cells[12] = valor
        return cells

    grid = WorkbookGrid(
        path="ipi.xls",
        sheet_name="Demonst. IPI 08-2026",
        kind="xlsx",
        rows=[
            ["DEMONSTRATIVO DO IPI"],
            ["APURAÇÃO"],
            row("Saldo credor do período anterior", 0),
            row("Total de débitos", 323063.81),
            row("Total de créditos", 324978.79),
            row("Saldo devedor de IPI", 0),
            row("Saldo credor de IPI para o mês seguinte", 1914.98),
        ],
    )
    parsed = parse_demonstrativo_ipi(grid)
    assert parsed["aRecolher"] == pytest.approx(-1914.98, abs=0.02)
    assert parsed["saldoCredorSeguinte"] == pytest.approx(1914.98, abs=0.02)
    patch = apuracao_patch_from_demo("ipi", parsed)
    assert patch["apuracao"]["ipi"]["aRecolher"] == pytest.approx(-1914.98, abs=0.02)
    assert patch["apuracao"]["ipi"]["saldoCredor"] == pytest.approx(1914.98, abs=0.02)


def test_jpg_demonstrativo_ipi_saldo_devedor():
    from app.extract.parse_impostos import parse_demonstrativo_ipi
    from app.extract.workbook import WorkbookGrid

    def row(label, valor):
        cells = [""] * 13
        cells[0] = label
        cells[12] = valor
        return cells

    grid = WorkbookGrid(
        path="ipi.xls",
        sheet_name="Demonst. IPI 01-2026",
        kind="xlsx",
        rows=[
            ["DEMONSTRATIVO DO IPI"],
            ["APURAÇÃO"],
            row("Total de débitos", 8225.69),
            row("Total de créditos", 1139.49),
            row("Saldo devedor de IPI", 7086.2),
            row("Saldo credor de IPI para o mês seguinte", 0),
        ],
    )
    parsed = parse_demonstrativo_ipi(grid)
    assert parsed["aRecolher"] == pytest.approx(7086.2, abs=0.02)


def test_jpg_ipi_saldo_credor_when_a_recolher_zero(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Empresa",
            "Filial",
            "Mês",
            "ICMS Crédito",
            "ICMS Débito",
            "ICMS a Recolher",
            "IPI Crédito",
            "IPI Débito",
            "IPI a Recolher",
        ]
    )
    ws.append(["JPG", "Filial PR", "Março", 0, 0, 0, 100, 20, 0])
    path = tmp_path / "ipi filial pr 01 a 08.xlsx"
    wb.save(path)
    result = classify_and_extract(path)
    mar = (result.get("parts") or [])[0]
    assert mar["pack_patch"]["apuracao"]["ipi"]["aRecolher"] == pytest.approx(-80, abs=0.02)
    assert mar["pack_patch"]["apuracao"]["ipi"]["credito"] == pytest.approx(100, abs=0.02)


def test_split_and_probe_movimento(tmp_path: Path):
    from scripts.split_movimento_mensal import split_file

    wb = Workbook()
    ws = wb.active
    ws.title = "Entradas"
    ws.append(["JPG PRODUTOS FUNCIONAIS"])
    ws.append(["CNPJ:", "21.051.983/0001-65"])
    ws.append(["Período: 01/01/2026 até 28/02/2026"])
    ws.append(
        [
            "Código",
            "Data Emissão",
            "Nota",
            "Série",
            "Esp",
            "X",
            "Y",
            "Z",
            "W",
            "Q",
            "Fornecedor",
            "A",
            "CNPJ/CPF",
            "IE",
            "B",
            "C",
            "CFOP",
            "AC",
            "UF",
            "Valor Contábil",
        ]
    )
    ws.append(
        [1, "15/01/2026", "100", "1", "", "", "", "", "", "", "FORN A", "", "11111111000111", "", "", "", "2.102", "", "GO", "100,00"]
    )
    ws.append(
        [2, "10/02/2026", "101", "1", "", "", "", "", "", "", "FORN B", "", "11111111000111", "", "", "", "2.102", "", "GO", "50,50"]
    )
    src = tmp_path / "Entradas 01-2026 a 02-2026.xlsx"
    wb.save(src)
    out = tmp_path / "out"
    reports = split_file(src, out, "entradas")
    assert {r["competencia"] for r in reports} == {"2026-01", "2026-02"}
    jan = next(p for p in out.glob("*.xlsx") if "01-2026" in p.name)
    result = classify_and_extract(jan)
    assert result["tipo"] == "entradas"
    assert result["company_id"] == "jpg"
    assert result["unidade"] == "sede"
    assert result["competencia"] == "2026-01"
    assert not result["errors"]
    assert abs(result["meta"]["delta"]) < 0.02
    assert result["pack_patch"]["totalCompras"] == pytest.approx(100, abs=0.02)


def test_aggregate_todas_does_not_invent():
    packs = [
        {"hasMovimentacao": True, "totalCompras": 10, "cfopSaidasTotal": 20, "nfsEntradas": 1, "nfsSaidas": 2},
        {"hasMovimentacao": True, "totalCompras": 3, "cfopSaidasTotal": 4, "nfsEntradas": 1, "nfsSaidas": 1},
    ]
    pack = aggregate_fiscal_packs(packs, "Jan/2026")
    assert pack["totalCompras"] == 13
    assert pack["cfopSaidasTotal"] == 24
    assert pack["nfsEntradas"] == 2


def test_aggregate_includes_irpj_csll():
    packs = [
        {
            "apuracao": {"irpj": {"aRecolher": 143211.70}, "csll": {"aRecolher": 74476.66}},
            "memoriaIrpj": {"aRecolher": 143211.70, "linhas": [{"label": "Saldo devedor de IRPJ", "valor": 143211.70}]},
            "memoriaCsll": {"aRecolher": 74476.66, "linhas": [{"label": "Saldo devedor", "valor": 74476.66}]},
        },
        {"apuracao": {"ipi": {"aRecolher": 40}}},
    ]
    pack = aggregate_fiscal_packs(packs, "1º Trimestre 2026")
    assert pack["apuracao"]["irpj"]["aRecolher"] == pytest.approx(143211.70, abs=0.02)
    assert pack["apuracao"]["csll"]["aRecolher"] == pytest.approx(74476.66, abs=0.02)
    assert pack["apuracao"]["ipi"]["aRecolher"] == pytest.approx(40, abs=0.02)
    assert pack["memoriaIrpj"]["aRecolher"] == pytest.approx(143211.70, abs=0.02)
    assert pack["memoriaCsll"]["aRecolher"] == pytest.approx(74476.66, abs=0.02)


FIXTURE_IRPJ_1T = ROOT / "fixtures" / "jpg-padrao" / "irpj-csll-1t-2026.xls"
FIXTURE_IRPJ_2T = ROOT / "fixtures" / "jpg-padrao" / "irpj-csll-2t-2026.xls"


@pytest.mark.skipif(not FIXTURE_IRPJ_1T.exists(), reason="Fixture IRPJ/CSLL 1T ausente")
def test_jpg_sede_irpj_csll_1t():
    result = classify_and_extract(FIXTURE_IRPJ_1T)
    assert result["tipo"] == "irpj_csll"
    assert result["company_id"] == "jpg"
    assert result["unidade"] == "sede"
    assert result["competencia"] == "2026-03"
    assert not result["errors"]
    ap = (result.get("pack_patch") or {}).get("apuracao") or {}
    assert ap["irpj"]["aRecolher"] == pytest.approx(143211.70, abs=0.02)
    assert ap["csll"]["aRecolher"] == pytest.approx(74476.66, abs=0.02)
    assert ap["irpj"]["fonte"] == "demonstrativo_exito_irpj"
    assert ap["csll"]["fonte"] == "demonstrativo_exito_csll"
    assert (result.get("pack_patch") or {}).get("memoriaIrpj", {}).get("aRecolher") == pytest.approx(143211.70, abs=0.02)


@pytest.mark.skipif(not FIXTURE_IRPJ_2T.exists(), reason="Fixture IRPJ/CSLL 2T ausente")
def test_jpg_sede_irpj_csll_2t():
    result = classify_and_extract(FIXTURE_IRPJ_2T)
    assert result["tipo"] == "irpj_csll"
    assert result["company_id"] == "jpg"
    assert result["unidade"] == "sede"
    assert result["competencia"] == "2026-06"
    assert not result["errors"]
    ap = (result.get("pack_patch") or {}).get("apuracao") or {}
    assert ap["irpj"]["aRecolher"] == pytest.approx(137737.02, abs=0.02)
    assert ap["csll"]["aRecolher"] == pytest.approx(77617.99, abs=0.02)


@pytest.mark.skipif(not FIXTURE_ENT.exists(), reason="Fixture jpg-padrao ausente")
def test_jpg_sede_jan_split_golden():
    path = FIXTURE_ENT
    result = classify_and_extract(path)
    assert result["company_id"] == "jpg"
    assert result["unidade"] == "sede"
    assert result["tipo"] == "entradas"
    assert not result["errors"]
    assert abs(result["meta"]["delta"]) < 0.02
